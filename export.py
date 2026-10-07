import os
import re
import sys
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from supabase import create_client

PAGE = 1000  # tope de filas por request de Supabase

# Colores de marca FlowPass (espejo de FlowPass-new/src/app.css)
VERDE, VERDE_SUAVE, OSCURO, GRIS, BORDE_GRIS = "00F28B", "E8FFF2", "13131A", "F5F5F5", "CBD5E1"
BORDE = Border(*(Side(style="thin", color=BORDE_GRIS),) * 4)

MEDIOS = {
    "cash": "Efectivo", "transfer": "Transferencia",
    "card-credit": "Tarjeta de crédito", "card-debit": "Tarjeta de débito",
    "yape": "Yape", "plin": "Plin", "codi": "CoDi",
}

load_dotenv()
url, key = os.getenv("SUPABASE_URL"), os.getenv("SUPABASE_SERVICE_ROLE_KEY")
gym_ids = [g.strip() for g in os.getenv("GYM_ID", "").split(",") if g.strip()]
if not url or not key or not gym_ids:
    sys.exit("Falta SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY o GYM_ID en .env")

sb = create_client(url, key)


def fetch_all(table, select, gym_id, **filters):
    """Trae todas las filas del gym paginando de a PAGE."""
    rows, start = [], 0
    while True:
        q = sb.table(table).select(select).eq("gym_id", gym_id)
        for col in filters.get("is_null", []):
            q = q.is_(col, "null")
        for col, val in filters.get("neq", {}).items():
            q = q.neq(col, val)
        batch = q.order("id").range(start, start + PAGE - 1).execute().data
        rows += batch
        if len(batch) < PAGE:
            return rows
        start += PAGE


def to_date(s):
    return date.fromisoformat(s[:10]) if s else None


def by_alumno(rows, fecha_col):
    # alumno (apellidos, nombre) y luego por fecha
    return sorted(rows, key=lambda r: ((r[2] or "").lower(), (r[3] or "").lower(), r[fecha_col] or date.min))


def alumnos(gym_id):
    rows = fetch_all(
        "members",
        "id,apellidos,nombre,codigo_ingreso,tipo_documento,identificacion,email,telefono,"
        "cumpleanos,direccion,fecha_alta,members_groups(groups(title))",
        gym_id, is_null=["deleted_at"],
    )
    out = [
        (m["id"], m["apellidos"], m["nombre"], m["codigo_ingreso"], m["tipo_documento"],
         m["identificacion"], m["email"], m["telefono"], to_date(m["cumpleanos"]),
         m["direccion"], to_date(m["fecha_alta"]),
         ", ".join(sorted(mg["groups"]["title"] for mg in m["members_groups"] if mg["groups"])))
        for m in rows
    ]
    out.sort(key=lambda r: ((r[1] or "").lower(), (r[2] or "").lower()))
    headers = ["id_alumno", "apellidos", "nombre", "codigo_ingreso", "tipo_documento",
               "identificacion", "email", "telefono", "cumpleanos", "direccion", "fecha_alta", "grupos"]
    return headers, out


def catalogo(gym_id):
    # incluye eliminados: compras y asistencias viejas los referencian
    rows = fetch_all(
        "planes",
        "id,label,amount,dias_mes,ilimitado,limite_clases,partners,is_free,neverExpires,order,deleted_at",
        gym_id,
    )
    rows.sort(key=lambda p: (p["deleted_at"] is not None, p["order"] or 0, (p["label"] or "").lower()))
    out = [
        (p["id"], (p["label"] or "").strip(), p["amount"],
         None if p["neverExpires"] else p["dias_mes"],
         "Ilimitado" if p["ilimitado"] else p["limite_clases"],
         "Sí" if p["partners"] else "No", "Sí" if p["is_free"] else "No",
         "No" if p["deleted_at"] else "Sí")
        for p in rows
    ]
    headers = ["id_plan", "paquete", "precio", "duracion_dias", "limite_clases",
               "compartido", "gratis", "activo"]
    return headers, out


def paquetes(gym_id):
    compras = fetch_all(
        "historico",
        "id,member_id,pago_id,plan_id,plan,fecha_inicio_plan,proxima_fecha_pago,monto,estado_pago,"
        "members!inner(apellidos,nombre),planes(label)",
        gym_id, is_null=["deleted_at", "members.deleted_at"],
    )
    # charges no tiene FK con historico: se cruzan por pago_id
    cargos = {
        c["pago_id"]: c
        for c in fetch_all(
            "charges",
            "id,pago_id,monto_total,monto_pendiente,estado,payments(monto_total,medio_de_pago,paid_at,estado)",
            gym_id, is_null=["deleted_at"], neq={"estado": "cancelado"},
        )
    }
    out = []
    for h in compras:
        c = cargos.get(h["pago_id"])
        pagos = [p for p in (c["payments"] if c else []) if p["estado"] == "confirmed"]
        out.append((
            h["member_id"], h["pago_id"], h["members"]["apellidos"], h["members"]["nombre"],
            h["plan_id"],
            (h["planes"] or {}).get("label") or h["plan"],
            to_date(h["fecha_inicio_plan"]),
            to_date(h["proxima_fecha_pago"]),
            c["monto_total"] if c else h["monto"],
            sum(p["monto_total"] for p in pagos) if c else None,
            c["monto_pendiente"] if c else None,
            c["estado"] if c else (h["estado_pago"] or "").lower() or None,
            ", ".join(sorted({MEDIOS.get(p["medio_de_pago"], p["medio_de_pago"]) for p in pagos if p["medio_de_pago"]})) or None,
            max((to_date(p["paid_at"]) for p in pagos if p["paid_at"]), default=None),
        ))
    headers = ["id_alumno", "id_compra", "apellidos", "nombre", "id_plan", "paquete", "inicio", "vence",
               "precio", "pagado", "debe", "estado_pago", "medio_de_pago", "fecha_ultimo_pago"]
    return headers, by_alumno(out, 6)


def asistencias(gym_id, tz):
    rows = fetch_all(
        "ingresos",
        "id,member_id,pago_id,plan_id,check_in,clases_tomadas,limite_clases,nota,"
        "members!inner(apellidos,nombre),planes(label),member_plans(plan),groups(title)",
        gym_id, is_null=["deleted_at", "members.deleted_at"], neq={"tipo": "UPDATE"},
    )
    out = [
        (i["member_id"], i["pago_id"], i["members"]["apellidos"], i["members"]["nombre"],
         # hora local del gym; Excel no guarda zona horaria
         datetime.fromisoformat(i["check_in"]).astimezone(tz).replace(tzinfo=None) if i["check_in"] else None,
         i["plan_id"],
         (i["planes"] or {}).get("label") or (i["member_plans"] or {}).get("plan"),
         (i["groups"] or {}).get("title"),
         i["clases_tomadas"], i["limite_clases"], i["nota"])
        for i in rows
    ]
    headers = ["id_alumno", "id_compra", "apellidos", "nombre", "fecha_hora", "id_plan", "paquete", "grupo",
               "clases_tomadas", "limite_clases", "nota"]
    return headers, by_alumno(out, 4)


def add_sheet(book, title, headers, rows):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for r in rows:
        sheet.append(r)
    for c in sheet[1]:
        c.font = Font(bold=True, color=OSCURO)
        c.fill = PatternFill("solid", fgColor=VERDE)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 20


# Pestaña "Cómo usar": lo que se le explica al cliente para cruzar pestañas.
GUIA = [
    ("CÓMO USAR ESTE ARCHIVO",),
    ["Creado por el equipo de desarrollo de FlowPass"],
    [],
    ("PESTAÑAS",),
    ["Alumnos", "Un alumno por fila."],
    ["Catálogo de paquetes", "Cada paquete que vende o vendió el negocio (activo = No: ya no se vende)."],
    ["Paquetes y pagos", "Cada paquete que compró o renovó un alumno, con lo pagado y lo que debe."],
    ["Asistencias", "Cada asistencia de un alumno, con la compra con la que asistió."],
    [],
    ("CÓMO SE RELACIONAN",),
    ("Código", "Qué identifica", "Está en", "Sirve para"),
    ["id_alumno", "Un alumno", "Alumnos, Paquetes y pagos, Asistencias",
     "Traer los datos del alumno de una compra o asistencia"],
    ["id_plan", "Un tipo de paquete", "Catálogo de paquetes, Paquetes y pagos, Asistencias",
     "Traer precio, duración y límite de clases del paquete"],
    ["id_compra", "Una compra concreta", "Paquetes y pagos, Asistencias",
     "Ver con qué compra asistió el alumno, o cuántas asistencias tuvo cada compra"],
    [],
    ("EJEMPLO DE FÓRMULA",),
    ["Precio del paquete en Asistencias (columna F = id_plan):"],
    ["=BUSCARV(F2; 'Catálogo de paquetes'!A:C; 3; FALSO)"],
    ["En Excel en inglés: =VLOOKUP(F2, 'Catálogo de paquetes'!A:C, 3, FALSE)"],
    [],
    ("NOTAS",),
    ["Si pagado y debe están vacíos, es un paquete antiguo: el estado viene en estado_pago."],
    ["clases_tomadas y limite_clases en Asistencias son los valores al momento de esa asistencia."],
    ["Las fechas y horas están en la zona horaria del negocio."],
    [],
    ("EJEMPLOS (los nombres y montos de aquí abajo son inventados)",),
    [],
    ("id_plan vs id_compra",),
    ["id_plan = QUÉ paquete es (el del catálogo). Es el mismo para todos los que compraron ese paquete."],
    ["id_compra = CUÁL compra fue. Cada compra o renovación tiene uno distinto."],
    ("Alumno", "Paquete", "id_plan", "id_compra", "Qué significa"),
    ["Juan", "Plan Mensual", "aaa", "111", "Juan compró el Plan Mensual en enero"],
    ["Juan", "Plan Mensual", "aaa", "222", "Juan lo renovó en febrero: mismo id_plan, otra id_compra"],
    ["María", "Plan Mensual", "aaa", "333", "María compró el mismo paquete que Juan: mismo id_plan"],
    ["María", "Plan Ilimitado", "bbb", "444", "Otro paquete: otro id_plan"],
    [],
    ("Cómo leer una fila de \"Paquetes y pagos\"",),
    ("Columna", "Valor", "Cómo se lee"),
    ["paquete", "Plan Mensual", "Qué compró"],
    ["inicio / vence", "01/02 – 01/03", "Desde y hasta cuándo puede usarlo"],
    ["precio", "150", "Lo que costaba"],
    ["pagado", "100", "Lo que ya pagó"],
    ["debe", "50", "Lo que le falta pagar"],
    ["estado_pago", "parcial", "pagado = completo, parcial = pagó una parte, pendiente = no pagó nada"],
    ["→ Se lee:", "Juan compró el Plan Mensual del 01/02 al 01/03, costaba 150, pagó 100 y debe 50."],
    [],
    ("Cómo leer una fila de \"Asistencias\"",),
    ("Columna", "Valor", "Cómo se lee"),
    ["fecha_hora", "05/02 18:30", "Cuándo asistió (hora del negocio)"],
    ["id_compra", "222", "Con qué compra asistió (la renovación de febrero)"],
    ["clases_tomadas / limite_clases", "3 / 8", "Era su clase 3 de 8 en ese momento (0 = ilimitado)"],
    ["→ Se lee:", "Juan asistió el 05/02 a las 18:30 con su Plan Mensual de febrero; iba en su clase 3 de 8."],
    [],
    ("Preguntas típicas y cómo responderlas",),
    ("Pregunta", "Dónde", "Cómo"),
    ["¿Qué paquetes compró un alumno?", "Paquetes y pagos", "Filtrar por id_alumno (o por apellidos)"],
    ["¿Cuánto debe un alumno?", "Paquetes y pagos", "Filtrar por id_alumno y sumar la columna debe"],
    ["¿Cuántas clases usó de una compra?", "Asistencias", "Filtrar por id_compra y contar las filas"],
    ["¿Cuánto cuesta un paquete?", "Catálogo de paquetes", "Buscar el id_plan"],
    ["¿Quién compró cierto paquete?", "Paquetes y pagos", "Filtrar por id_plan (o por paquete)"],
]


def add_guide(book, title, rows, widths):
    # primera fila = título, segunda = firma; tupla de 1 = sección; tupla de varios = encabezado de tabla
    # (las filas que siguen a un encabezado, hasta la próxima vacía, van con borde)
    sheet = book.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    in_table = False
    for i, row in enumerate(rows):
        sheet.append(list(row))
        n = sheet.max_row
        cells = [sheet.cell(n, col) for col in range(1, len(widths) + 1)]
        for c in cells:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if not row:
            in_table = False
        elif i == 0:
            cells[0].font = Font(bold=True, size=18, color=OSCURO)
        elif i == 1:
            cells[0].font = Font(italic=True, color="64748B")
        elif isinstance(row, tuple) and len(row) == 1:
            for c in cells:
                c.font = Font(bold=True, size=12, color=VERDE)
                c.fill = PatternFill("solid", fgColor=OSCURO)
        elif isinstance(row, tuple):
            in_table = True
            for c in cells[:len(row)]:
                c.font = Font(bold=True, color=OSCURO)
                c.fill = PatternFill("solid", fgColor=VERDE_SUAVE)
                c.border = BORDE
        elif in_table:
            for c in cells[:len(row)]:
                c.border = BORDE
        if row and row[0].startswith("="):
            cells[0].data_type = "s"  # mostrar la fórmula de ejemplo, no calcularla
            cells[0].font = Font(name="Consolas")
            cells[0].fill = PatternFill("solid", fgColor=GRIS)
        if row and row[0].startswith("→"):
            for c in cells:
                c.font = Font(italic=True)
    for letter, width in zip("ABCDE", widths):
        sheet.column_dimensions[letter].width = width


Path("output").mkdir(exist_ok=True)
for gym_id in gym_ids:
    gym = sb.table("gyms").select("name,timezone").eq("id", gym_id).maybe_single().execute()
    if not gym or not gym.data:
        print(f"GYM_ID {gym_id} no existe, lo salto")
        continue
    name, tz = gym.data["name"], ZoneInfo(gym.data["timezone"] or "America/Lima")

    book = Workbook()
    book.remove(book.active)
    add_guide(book, "Cómo usar", GUIA, (32, 22, 45, 20, 60))
    for title, (headers, rows) in [
        ("Alumnos", alumnos(gym_id)),
        ("Catálogo de paquetes", catalogo(gym_id)),
        ("Paquetes y pagos", paquetes(gym_id)),
        ("Asistencias", asistencias(gym_id, tz)),
    ]:
        add_sheet(book, title, headers, rows)
        print(f"{name} · {title}: {len(rows)} filas")

    out = Path("output") / f"{re.sub(r'[^\w\- ]', '', name).strip()}.xlsx"
    book.save(out)
    print(f"→ {out}")
